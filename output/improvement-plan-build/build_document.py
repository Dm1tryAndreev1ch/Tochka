"""Build the editable project plan and its architecture diagram from Markdown."""
from pathlib import Path
import json
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/PROJECT_IMPROVEMENT_PLAN.md"
OUT = ROOT / "output/tochka-improvement-and-security-plan-2026-10-13.docx"
FIG = ROOT / "docs/architecture/tochka-security-target.png"
TOC = Path(__file__).with_name("toc.json")


def diagram():
    image = Image.new("RGB", (1800, 1510), "white")
    d = ImageDraw.Draw(image)
    fonts = "/System/Library/Fonts/Supplemental/"
    normal = ImageFont.truetype(fonts + "Arial.ttf", 36)
    bold = ImageFont.truetype(fonts + "Arial Bold.ttf", 38)
    small = ImageFont.truetype(fonts + "Arial.ttf", 31)
    title = ImageFont.truetype(fonts + "Arial Bold.ttf", 47)
    d.text((55, 28), "Границы доверия и основные связи", font=title, fill="#17222d")
    d.text((55, 91), "Каждый блок — отдельные identity, права, секреты и сетевые правила", font=small, fill="#505a64")

    def box(rect, heading, lines, color="#f4f6f8"):
        x, y, x2, y2 = rect
        d.rounded_rectangle(rect, radius=12, fill=color, outline="#aab4be", width=2)
        d.text((x+25, y+20), heading, font=bold, fill="#162635")
        for i, line in enumerate(lines):
            d.text((x+25, y+79+i*46), line, font=normal, fill="#243645")

    def down(x, y, end):
        d.line((x, y, x, end-15), fill="#506d85", width=5)
        d.polygon([(x-12, end-18), (x+12, end-18), (x, end)], fill="#506d85")

    box((55, 155, 1105, 365), "01  Вход и медиа", ["Edge / операторский доступ / WebRTC", "SIP и Telegram — отдельные адаптеры"])
    down(580, 365, 423)
    box((55, 425, 1105, 635), "02  Приложение и бизнес", ["Session / Dialogue / Action / Knowledge", "Права, подтверждения, транзакция и outbox"], "#edf4fa")
    down(310, 635, 705)
    down(850, 635, 705)
    box((55, 710, 565, 998), "03  Вычисления", ["ASR / LLM / TTS", "Mac или Windows GPU", "Без БД, KMS и CRM", "Без выхода в интернет"])
    box((595, 710, 1105, 998), "04  Данные", ["Business DB / PII Vault", "Архив: шифрование", "Ключи через broker", "Свои роли и storage"])
    box((55, 1090, 1105, 1300), "05  Интеграции", ["CRM / календарь / внешние провайдеры", "Узкий API outbox; egress по разрешённым адресам"])
    d.line((1080, 635, 1142, 635, 1142, 1055, 580, 1055, 580, 1074), fill="#506d85", width=4)
    d.polygon([(568,1072),(592,1072),(580,1090)],fill="#506d85")

    box((1200, 155, 1745, 665), "06  Безопасность", ["PKI / identity / policy", "KMS/HSM: свои роли", "Audit Ledger + signer", "Независимый witness", "Verifier + deletion ledger", "Без изменения бизнеса"], "#edf4fa")
    box((1200, 725, 1745, 1010), "07  Управление", ["Bastion / CI / registry", "Backup и recovery", "Отдельные администраторы", "Двойной контроль"])
    box((1200, 1090, 1745, 1300), "Наблюдаемость", ["Метрики и сигналы", "Без текста и аудио"], "#f7f7f7")
    d.line((1107, 530, 1180, 530), fill="#506d85", width=4)
    d.polygon([(1179,518),(1179,542),(1197,530)], fill="#506d85")
    d.text((55, 1360), "Стрелки показывают основные связи; полный allowlist — в разделе 7.", font=normal, fill="#374b5d")
    d.text((55, 1420), "В промышленном профиле данные и ключи разделены также VM/узлами и полномочиями.", font=small, fill="#505a64")
    FIG.parent.mkdir(exist_ok=True)
    image.save(FIG)


def set_cell_shading(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), fill)
    pr.append(el)


def hyperlink(p, text, target=None, anchor=None):
    h = OxmlElement("w:hyperlink")
    if target:
        h.set(qn("r:id"), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    if anchor:
        h.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "174C75" if target else "000000")
    pr.append(color)
    r.append(pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


def inline(p, text):
    pattern = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")
    for s in pattern.split(text):
        if not s:
            continue
        if s.startswith("[") and "](" in s:
            m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", s)
            label, href = m.groups()
            if not href.startswith(("https://", "http://")):
                href = str((SOURCE.parent / href).resolve())
            hyperlink(p, label, href)
        elif s.startswith("**"):
            p.add_run(s[2:-2]).bold = True
        elif s.startswith("`"):
            p.add_run(s[1:-1])
        else:
            p.add_run(s)


def bookmark(p, name, num):
    el = OxmlElement("w:bookmarkStart")
    el.set(qn("w:id"), str(num))
    el.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(num))
    p._p.insert(0, el)
    p._p.append(end)


def field(p, instruction):
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), instruction)
    p._p.append(f)


def build():
    diagram()
    lines = SOURCE.read_text().splitlines()
    doc = Document()
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.72)
    sec.left_margin = sec.right_margin = Inches(.72)
    sec.header_distance = sec.footer_distance = Inches(.3)
    sec.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0,0,0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.widow_control = True
    for style, size in [("Title",28),("Subtitle",13),("Heading 1",19),("Heading 2",14),("Heading 3",12)]:
        s = doc.styles[style]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0,0,0)
        s.font.bold = style != "Subtitle"
        s.paragraph_format.space_before = Pt(17 if style.startswith("Heading") else 0)
        s.paragraph_format.space_after = Pt(8)
        s.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.size = Pt(11)
    header = sec.header.paragraphs[0]
    header.text = "ТОЧКА  ·  ПРОГРАММА РАЗВИТИЯ И ЗАЩИТЫ"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("666666")
    footer = sec.footer.paragraphs[0]
    doc.styles["Footer"].font.name = "Arial"
    doc.styles["Footer"].font.size = Pt(8)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Версия 1.0  ·  13.10.2026     ").font.size = Pt(8)
    field(footer, "PAGE")
    footer.add_run(" / ").font.size = Pt(8)
    field(footer, "NUMPAGES")

    first_section = next(i for i,x in enumerate(lines) if x.startswith("## "))
    doc.add_paragraph("ТОЧКА", "Subtitle")
    p=doc.add_paragraph(lines[0][2:], "Title")
    p.paragraph_format.space_before = Pt(38)
    for line in lines[1:first_section]:
        if line.strip():
            p=doc.add_paragraph()
            inline(p,line)
    p=doc.add_paragraph("Профили: локальный Mac · Mac + Windows RTX 5070 Mobile · усиленное промышленное размещение")
    p.paragraph_format.space_before=Pt(20)
    p.runs[0].bold=True
    doc.add_page_break()
    doc.add_paragraph("Содержание", "Heading 1")
    titles=[x[3:] for x in lines if x.startswith("## ")]
    toc_pages=json.loads(TOC.read_text()) if TOC.exists() else {}
    for i,title in enumerate(titles):
        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.line_spacing=1
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.95), WD_ALIGN_PARAGRAPH.RIGHT)
        hyperlink(p,title,anchor=f"section_{i}")
        p.add_run("\t"+str(toc_pages.get(title,"")))
    doc.add_page_break()

    section_index=0
    i=first_section
    while i<len(lines):
        text=lines[i]
        if not text.strip():
            i+=1; continue
        if text.startswith("#"):
            m=re.match(r"(#+) (.+)",text)
            level=len(m.group(1))-1
            p=doc.add_paragraph(m.group(2), f"Heading {level}")
            if level==1:
                bookmark(p, f"section_{section_index}",section_index+1)
                section_index+=1
            i+=1; continue
        if text.startswith("!["):
            m=re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)",text)
            p=doc.add_paragraph()
            p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(SOURCE.parent/m.group(2)),width=Inches(7.02))
            p=doc.add_paragraph(m.group(1))
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size=Pt(9)
            i+=1;continue
        if text.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                cells=[x.strip() for x in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",x) for x in cells):
                    rows.append(cells)
                i+=1
            n=len(rows[0])
            table=doc.add_table(rows=0,cols=n)
            table.autofit=False
            ratios=([.16,.35,.49] if rows[0][0]=="ID" else [.24,.39,.37]) if n==3 else ([.20,.29,.27,.24] if n==4 else [1/n]*n)
            total=7.02
            for col,w in zip(table.columns,ratios):col.width=Inches(total*w)
            props=table._tbl.tblPr
            borders=OxmlElement("w:tblBorders")
            for side in ["top","left","bottom","right","insideH","insideV"]:
                b=OxmlElement("w:"+side)
                b.set(qn("w:val"),"single");b.set(qn("w:sz"),"4");b.set(qn("w:color"),"D9D9D9")
                borders.append(b)
            props.append(borders)
            margins=OxmlElement("w:tblCellMar")
            for side, val in [("top","85"),("bottom","85"),("left","95"),("right","95")]:
                el=OxmlElement("w:"+side);el.set(qn("w:w"),val);el.set(qn("w:type"),"dxa");margins.append(el)
            props.append(margins)
            for rowno,values in enumerate(rows):
                row=table.add_row()
                trpr=row._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if rowno==0:trpr.append(OxmlElement("w:tblHeader"))
                for colno,(cell,value) in enumerate(zip(row.cells,values)):
                    cell.width=Inches(total*ratios[colno])
                    p=cell.paragraphs[0]
                    p.paragraph_format.space_after=Pt(1)
                    p.paragraph_format.line_spacing=1.02
                    inline(p,value)
                    for run in p.runs:
                        run.font.size=Pt(10.5)
                        if rowno==0:run.bold=True
                    if rowno==0:set_cell_shading(cell,"F2F2F2")
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        p=doc.add_paragraph(style="List Bullet" if text.startswith("- ") else None)
        inline(p,text[2:] if text.startswith("- ") else text)
        i+=1
    doc.core_properties.title="План совершенствования и защиты голосовой платформы Точка"
    doc.core_properties.subject="Качество голоса, модульная архитектура, ИБ и распределённый inference"
    doc.core_properties.author="Точка"
    doc.core_properties.language="ru-RU"
    settings=doc.settings.element
    update=OxmlElement("w:updateFields");update.set(qn("w:val"),"true");settings.append(update)
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(json.dumps({"docx":str(OUT),"sections":section_index,"words":len(SOURCE.read_text().split()),"bytes":OUT.stat().st_size},ensure_ascii=False))


if __name__=="__main__":build()
